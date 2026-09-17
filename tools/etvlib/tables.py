"""A table over RPM and one further axis, and reading and writing it as CSV or Excel.

All three tables of the app have this form: rows are RPM, columns are
throttle % (engine torque) or grip % (torque request, ETV map).

The file layout is the one a spreadsheet gives: the first row holds the
column breakpoints, the first column the RPM, the corner cell any label.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass

import numpy as np


class TableError(ValueError):
    """A file or a text that is not a table. The message is for the user."""


@dataclass(frozen=True)
class Table:
    """Values over RPM (rows) and one further axis (columns), both ascending."""

    rpm: np.ndarray
    axis: np.ndarray
    values: np.ndarray

    def __post_init__(self) -> None:
        if self.values.shape != (len(self.rpm), len(self.axis)):
            raise TableError(f"A table of {len(self.rpm)} × {len(self.axis)} breakpoints "
                             f"cannot hold values of shape {self.values.shape}.")


def make_table(rpm, axis, values) -> Table:
    """A table from breakpoints in any order — sorted ascending on both axes."""
    rpm = np.asarray(rpm, dtype=float)
    axis = np.asarray(axis, dtype=float)
    values = np.asarray(values, dtype=float)
    if values.ndim != 2 or values.shape != (len(rpm), len(axis)):
        raise TableError(f"{len(rpm)} RPM and {len(axis)} column breakpoints do not fit "
                         f"values of shape {values.shape}.")
    rows, cols = np.argsort(rpm, kind="stable"), np.argsort(axis, kind="stable")
    return Table(rpm[rows], axis[cols], values[np.ix_(rows, cols)])


def from_json(obj, what: str) -> Table:
    """A table as a page sends it: `{"rpm": […], "axis": […], "values": [[…], …]}`.

    `what` names it in the message ("engine torque table"). An editable grid
    can send anything — an empty cell, a breakpoint typed twice — so
    everything is checked here, once, for every route."""
    if not isinstance(obj, dict) or not all(key in obj for key in ("rpm", "axis", "values")):
        raise TableError(f"The {what} is missing.")
    try:
        table = make_table(obj["rpm"], obj["axis"], obj["values"])
    except (TypeError, ValueError) as error:
        if isinstance(error, TableError):
            raise TableError(f"The {what}: {error}") from None
        raise TableError(f"The {what} has a cell that is not a number.") from None
    if not len(table.rpm) or not len(table.axis):
        raise TableError(f"The {what} is empty.")
    if not (np.isfinite(table.rpm).all() and np.isfinite(table.axis).all() and np.isfinite(table.values).all()):
        raise TableError(f"The {what} has an empty cell or one that is not a number.")
    for name, breakpoints in (("RPM", table.rpm), ("column", table.axis)):
        twice = breakpoints[:-1][np.diff(breakpoints) == 0]
        if len(twice):
            raise TableError(f"The {what} has the {name} breakpoint {format_breakpoint(twice[0])} twice.")
    return table


def to_json(table: Table) -> dict:
    return {"rpm": table.rpm.tolist(), "axis": table.axis.tolist(), "values": table.values.tolist()}


def parse_breakpoints(text: str) -> np.ndarray:
    """Breakpoints typed as `0, 2, 3.5, …` — sorted ascending.

    Raises `ValueError` when a part is not a number."""
    return np.array(sorted(float(part) for part in text.split(",") if part.strip()))


def format_breakpoint(value: float) -> str:
    """`4000` rather than `4000.0`, `12.5` as it is."""
    value = float(value)
    return str(int(value)) if value == int(value) else str(value)


def _number(cell, where: str) -> float:
    if cell is None or (isinstance(cell, str) and not cell.strip()):
        raise TableError(f"{where} is empty.")
    try:
        return float(cell)
    except (TypeError, ValueError):
        raise TableError(f"{where} is not a number: {cell!r}.") from None


def _from_rows(rows: list[list]) -> Table:
    rows = [r for r in rows if any(c is not None and str(c).strip() for c in r)]
    if len(rows) < 2 or len(rows[0]) < 2:
        raise TableError("A table needs a row of column breakpoints and at least one row of values.")
    axis = [_number(c, f"Column breakpoint {i + 1}") for i, c in enumerate(rows[0][1:])]
    rpm, values = [], []
    for n, row in enumerate(rows[1:], start=2):
        if len(row) != len(axis) + 1:
            raise TableError(f"Row {n} has {len(row) - 1} values, the header has {len(axis)} breakpoints.")
        rpm.append(_number(row[0], f"The RPM of row {n}"))
        values.append([_number(c, f"Row {n}, column {i + 1}") for i, c in enumerate(row[1:])])
    return make_table(rpm, axis, values)


def read_csv(text: str) -> Table:
    """A table from CSV text. Also reads what a German spreadsheet exports:
    semicolons between cells and a decimal comma."""
    text = text.lstrip("\ufeff")
    first = text.split("\n", 1)[0]
    if ";" in first:
        rows = [[c.replace(",", ".") for c in row] for row in csv.reader(io.StringIO(text), delimiter=";")]
    else:
        rows = list(csv.reader(io.StringIO(text)))
    return _from_rows(rows)


def to_csv(table: Table, corner: str = "") -> str:
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\n")
    writer.writerow([corner, *(repr(float(v)) for v in table.axis)])
    for rpm, row in zip(table.rpm, table.values, strict=True):
        writer.writerow([repr(float(rpm)), *(repr(float(v)) for v in row)])
    return out.getvalue()


def read_xlsx(data: bytes) -> Table:
    """A table from the first sheet of an Excel file."""
    from openpyxl import load_workbook  # noqa: PLC0415 — only needed for Excel files

    sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
    return _from_rows([list(row) for row in sheet.iter_rows(values_only=True)])


def to_xlsx(table: Table, sheet_name: str, corner: str = "") -> bytes:
    from openpyxl import Workbook  # noqa: PLC0415 — only needed for Excel files

    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name[:31]
    sheet.append([corner, *(float(v) for v in table.axis)])
    for rpm, row in zip(table.rpm, table.values, strict=True):
        sheet.append([float(rpm), *(float(v) for v in row)])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
